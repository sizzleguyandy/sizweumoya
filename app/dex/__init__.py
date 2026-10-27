"""Dex, the run brain: writes running programmes. Only Dex writes them; Sarah never does.

Dex is code, not a model: the TR3D Run Coach engine (Daniels VO2X paces, phases, volume curve,
quality sessions, hills, taper) builds the plan, and Dex turns it into programme rows with
Coach Andy's heart-rate targets on every session. A plan is only ever a draft: it goes to
Coach Andy as an approval, and is loaded (and sent to the watch) only when he approves it.
"""
import json
import re
from datetime import date, timedelta

from .. import paces, programme, queries
from ..config import settings
from .engine.plan_builder import build_full_plan

MIN_WEEKS = 6
MAX_WEEKS = 24
DAYS = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]
PROFILES = ("conservative", "aggressive")
PHASES = {1: "Base", 2: "Repetitions", 3: "Intervals", 4: "Threshold and taper"}


class DexError(ValueError):
    pass


def race_distance(km: float) -> str:
    """Our race distance in km -> the engine's distance class."""
    for limit, name in ((6, "5k"), (12, "10k"), (25, "half"), (45, "marathon"), (60, "ultra_56")):
        if km <= limit:
            return name
    return "ultra_90"


def hilliness(distance_km: float, ascent_m: float | None) -> str:
    per_km = (ascent_m or 0) / distance_km if distance_km else 0
    return "low" if per_km < 5 else "medium" if per_km < 12 else "high"


def _next_monday(d: date) -> date:
    return d + timedelta(days=(7 - d.weekday()) % 7 or 7)


def _parse_days(text: str) -> list[str]:
    out = []
    for part in re.split(r"[,\s]+", text or ""):
        p = part.strip()[:3].capitalize()
        if p and p not in DAYS:
            raise DexError(f"{part!r} isn't a day (use Mon, Tue, ... Sun)")
        if p and p not in out:
            out.append(p)
    return out


# ---- what Dex knows about the runner ------------------------------------------------------

def inputs(conn, rid: int, *, race_name: str = "", weekly_km: float | None = None,
           vdot: float | None = None, today: date | None = None) -> dict:
    """Everything the plan is built from, taken from the record unless Coach Andy gives it."""
    today = today or date.today()
    races = conn.execute("SELECT name, date, distance_km, ascent_m FROM races WHERE runner_id=? AND date>?"
                         " ORDER BY date", (rid, today.isoformat())).fetchall()
    if race_name:
        races = [r for r in races if race_name.strip().lower() in r["name"].lower()]
    if not races:
        raise DexError("no upcoming race on file" + (f" matching {race_name!r}" if race_name else "")
                       + " (save one first: save_race or the Races sheet)")
    race = dict(races[0])
    base = queries.baseline(conn, rid, weeks=6, today=today)
    km = weekly_km if weekly_km is not None else base.get("avg_km_per_week")
    if not km:
        raise DexError("no recent runs to work out weekly km from; give weekly_km")
    if vdot is None:
        bench = paces.best_benchmark(conn, rid, today)
        if bench is None:
            raise DexError("no race or time trial on file to set paces from; give vdot, or have them run "
                           "a flat all-out 5 km first")
        vdot, vdot_from = bench["vdot"], f"{bench['distance_km']} km on {bench['date']}"
    else:
        vdot_from = "given by Coach Andy"
    return {"race": race, "weekly_km": round(float(km), 1), "vdot": round(float(vdot), 1), "vdot_from": vdot_from,
            "longest_recent_km": base.get("longest_run_km")}


# ---- heart-rate targets -------------------------------------------------------------------

def _zones(runner) -> list[int]:
    z = json.loads(runner["hr_zones"]) if runner["hr_zones"] else None
    if not z or len(z) != 4:
        raise DexError("this runner has no heart rate zones yet; set or approve them first")
    return [int(x) for x in z]


def classify(session: str, zones: list[int], max_hr: int | None) -> tuple[str, list[int] | None]:
    """(kind, hr_range) for one engine session. Easy and long runs get the runner's easy cap
    from their profile, so no range here."""
    s = session.lower()
    if "race day" in s:
        return "race", None
    if s.startswith("long run") or s.startswith("back-to-back"):
        return "long", None
    if s.startswith("threshold"):
        return "workout", [zones[2] + 1, zones[3]]
    if s.startswith("intervals") or s.startswith("long hill"):
        return "workout", [zones[3] + 1, max_hr or zones[3] + 15]
    if s.startswith(("repetitions", "short hill", "downhill")):
        return "workout", None            # too short for heart rate: run them by pace or effort
    return "easy", None                   # easy, recovery, medium-long, strides, run/walk


# ---- the plan --------------------------------------------------------------------------------

def draft(conn, rid: int, *, race_name: str = "", weekly_km: float | None = None, vdot: float | None = None,
          long_day: str = "Sun", quality_day: str = "Tue", extra_days: str = "Thu",
          profile: str = "conservative", start: date | None = None, today: date | None = None) -> dict:
    """Build a programme draft: {"inputs", "summary", "rows", "warnings"}. Writes nothing."""
    today = today or date.today()
    if profile not in PROFILES:
        raise DexError(f"profile must be one of {', '.join(PROFILES)}")
    long_ = _parse_days(long_day)
    quality = _parse_days(quality_day)
    extra = _parse_days(extra_days)
    if len(long_) != 1 or len(quality) != 1:
        raise DexError("give one long-run day and one quality day")
    if long_[0] == quality[0] or set(extra) & {long_[0], quality[0]}:
        raise DexError("the long run, quality and extra days must all be different days")
    runner = conn.execute("SELECT * FROM runners WHERE runner_id=?", (rid,)).fetchone()
    zones = _zones(runner)
    inp = inputs(conn, rid, race_name=race_name, weekly_km=weekly_km, vdot=vdot, today=today)
    race = inp["race"]
    race_day = date.fromisoformat(race["date"])
    race_monday = race_day - timedelta(days=race_day.weekday())
    first = start or _next_monday(today)
    first -= timedelta(days=first.weekday())                      # weeks run Monday to Sunday
    weeks = (race_monday - first).days // 7 + 1
    if weeks < MIN_WEEKS:
        raise DexError(f"only {max(weeks, 0)} weeks until {race['name']} ({race['date']}); Dex needs at least "
                       f"{MIN_WEEKS}. Write a short block like this one by hand (programme template)")
    if weeks > MAX_WEEKS:
        first = race_monday - timedelta(weeks=MAX_WEEKS - 1)
        weeks = MAX_WEEKS
    plan = build_full_plan(inp["weekly_km"], inp["vdot"], race_distance(race["distance_km"]), race_day, first,
                           race_hilliness=hilliness(race["distance_km"], race["ascent_m"]),
                           long_run_day=long_[0], quality_day=quality[0], training_profile=profile,
                           extra_training_days=",".join(extra) or "Thu", total_weeks=weeks)
    rows, warnings = [], []
    if (plan["base_building"] or {}).get("warning"):
        warnings.append(plan["base_building"]["warning"].lstrip("⚠ℹ\ufe0f "))
    for w in plan["weeks"]:
        monday = date.fromisoformat(w["week_start"])
        for name, s in w["days"].items():
            day = monday + timedelta(days=DAYS.index(name))
            kind, hr = classify(s["session"], zones, runner["max_hr"])
            if kind == "race":
                if day != race_day:
                    raise DexError(f"race day came out on {day}, not {race_day}; nothing was drafted")
                km = race["distance_km"]
            else:
                km = s.get("km") or 0
                if km <= 0:
                    continue                                       # rest, walk or cross-training
            label = race["name"] if kind == "race" else s["session"].replace("🏁 ", "")
            rows.append({"date": day.isoformat(), "type": "run", "kind": kind, "km": f"{km:g}",
                         "description": f"{label} — {s['notes']}"[:300],
                         "hr_low": str(hr[0]) if hr else "", "hr_high": str(hr[1]) if hr else "",
                         "climb_m": f"{race['ascent_m']:g}" if kind == "race" and race["ascent_m"] else "",
                         "week": str(w["week_number"]), "phase": PHASES.get(w["phase"], str(w["phase"]))})
    if inp["longest_recent_km"] and rows:
        first_long = next((float(r["km"]) for r in rows if r["kind"] == "long"), None)
        if first_long and first_long > 1.3 * inp["longest_recent_km"]:
            warnings.append(f"the first long run ({first_long:g} km) is well above their longest recent run "
                            f"({inp['longest_recent_km']:g} km)")
    vols = [w["planned_volume_km"] for w in plan["weeks"]]
    longs = [float(r["km"]) for r in rows if r["kind"] == "long"]
    summary = {
        "race": f"{race['name']} {race['distance_km']:g} km on {race['date']}",
        "weeks": weeks, "start": first.isoformat(),
        "phases": {PHASES[i]: plan["phases"][f"phase_{n}_weeks"]
                   for i, n in ((1, "I"), (2, "II"), (3, "III"), (4, "IV")) if plan["phases"][f"phase_{n}_weeks"]},
        "weekly_km": {"start": vols[0], "peak": max(vols), "race_week": vols[-1]},
        "longest_run_km": max(longs) if longs else None,
        "sessions": len(rows),
        "days": {"long": long_[0], "quality": quality[0], "extra": extra},
        "paces": {k: v for k, v in plan["paces"].items() if k != "vo2x"},
        "easy_hr_cap": runner["easy_hr_cap"], "profile": profile,
    }
    return {"inputs": inp, "summary": summary, "rows": rows, "warnings": warnings}


def summary_text(name: str, d: dict) -> str:
    s, i = d["summary"], d["inputs"]
    lines = [f"Dex programme draft for {name}: {s['race']}",
             f"{s['weeks']} weeks from {s['start']} ("
             + ", ".join(f"{k} {v}" for k, v in s["phases"].items()) + ")",
             f"Weekly km {s['weekly_km']['start']:g} -> peak {s['weekly_km']['peak']:g}; longest run "
             f"{s['longest_run_km']:g} km; {s['sessions']} sessions",
             f"Days: long {s['days']['long']}, quality {s['days']['quality']}, extra {', '.join(s['days']['extra'])}",
             f"Built from {i['weekly_km']:g} km/week and VDOT {i['vdot']:g} ({i['vdot_from']}); profile {s['profile']}",
             f"Easy pace {s['paces']['easy']}, threshold {s['paces']['threshold']}; easy HR cap {s['easy_hr_cap']}"]
    lines += [f"Note: {w}" for w in d["warnings"]]
    return "\n".join(lines)


def week_lines(d: dict, weeks: int = 2) -> list[str]:
    """The first weeks session by session, for a quick look in the chat."""
    out = []
    for r in d["rows"]:
        if int(r["week"]) > weeks:
            break
        hr = f" HR {r['hr_low']}-{r['hr_high']}" if r["hr_low"] else ""
        out.append(f"W{r['week']} {date.fromisoformat(r['date']):%a %d %b}  {r['km']} km {r['kind']}{hr}: "
                   f"{r['description'][:70]}")
    return out


# ---- the draft as a file Coach Andy can open, edit and load ---------------------------------

FILE_COLS = ("date", "type", "description", "km", "kind", "hr_low", "hr_high", "climb_m", "exercises", "steps")


def write_xlsx(d: dict, path: str) -> str:
    """The draft in the programme template (templates/programme_template.xlsx), so an edited
    copy loads with `load-programme --replace` like any other programme."""
    from pathlib import Path

    import openpyxl
    tpl = Path(__file__).resolve().parents[2] / "templates" / "programme_template.xlsx"
    wb = openpyxl.load_workbook(tpl)
    ws = wb["Programme"]
    head = [str(c.value or "").strip().lower() for c in ws[1]]
    for extra in ("week", "phase"):
        if extra not in head:
            ws.cell(row=1, column=len(head) + 1, value=extra)
            head.append(extra)
    for i, r in enumerate(d["rows"], start=2):
        for j, col in enumerate(head, start=1):
            v = r.get(col, "")
            if col in ("km", "hr_low", "hr_high", "climb_m", "week") and v not in ("", None):
                v = float(v) if "." in str(v) else int(v)
            ws.cell(row=i, column=j, value=v if v != "" else None)
    if "Example" in wb.sheetnames:
        del wb["Example"]
    wb.save(path)
    return path


# ---- draft -> Coach Andy -> programme ------------------------------------------------------

def drafts_dir():
    from pathlib import Path
    d = Path(settings.db_path).resolve().parent / "dex"
    d.mkdir(parents=True, exist_ok=True)
    return d


def propose(conn, rid: int, d: dict) -> dict:
    """Save the draft as a 'programme' approval for Coach Andy, with the file next to it."""
    r = conn.execute("SELECT name, hermes_profile FROM runners WHERE runner_id=?", (rid,)).fetchone()
    text = summary_text(r["name"], d)
    aid = queries.create_approval(conn, rid, "Dex programme draft", text, "programme",
                                  {"rows": d["rows"], "summary": d["summary"]})
    path = str(drafts_dir() / f"dex-{r['hermes_profile'] or rid}-{aid}.xlsx")
    write_xlsx(d, path)
    return {"approval_id": aid, "file": path, "text": text}


def apply(conn, approval: dict) -> dict | None:
    """Coach Andy approved a Dex draft: load it as the runner's run programme (replacing the
    current one). None if the approval isn't a Dex draft."""
    if approval.get("kind") != "programme" or not approval.get("payload"):
        return None
    rows = json.loads(approval["payload"])["rows"]
    try:
        out = programme.load_rows(conn, approval["runner_id"], rows, replace=True)["run"]
    except programme.ProgrammeError as e:
        return {"error": str(e)}
    return out
