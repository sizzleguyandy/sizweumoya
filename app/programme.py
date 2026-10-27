"""Load a programme Coach Andy wrote (run, strength, or both) from the Excel template
(templates/programme_template.xlsx, sheet "Programme", optional sheet "Races") or a CSV.

One row per session:

  date,type,description,km,kind,hr_low,hr_high,exercises,steps
  2026-11-03,run,Easy run,5,easy,,144,
  2026-11-05,run,Hill reps 6x90s,7,workout,160,175,,"- 15m Z2 HR|6x|- 90s Z4 HR|- 90s Z1 HR|- 10m Z2 HR"
  2026-11-04,strength,Lower body A,,,,,"Squat 4x6 70kg; RDL 3x8 60kg; Calf raise 3x15"

type: run or strength. kind (runs): easy, long, workout, benchmark, race. Empty cells are fine.
Optional column climb_m (target climb for a run). Races sheet: name, date, distance_km, ascent_m, note.
Loading a programme moves the person to "guided" for that discipline until its last date.
"""
import csv
import json
from datetime import date, datetime

RUN_KINDS = ("easy", "long", "workout", "benchmark", "race")


class ProgrammeError(ValueError):
    pass


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _sheet_rows(ws) -> list[dict]:
    it = ws.iter_rows(values_only=True)
    head = [(_cell(h) or "").lower() for h in next(it, [])]
    out = []
    for r in it:
        row = {head[i]: _cell(v) for i, v in enumerate(r) if i < len(head) and head[i]}
        if any(row.values()) and not row.get("date", "").lower().startswith(("e.g", "example")):
            out.append(row)
    return out


def read_rows(path: str) -> tuple[list[dict], list[dict]]:
    """(sessions, races) from the Excel template or a CSV of sessions."""
    if path.lower().endswith((".xlsx", ".xlsm")):
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True)
        if "Programme" not in wb.sheetnames:
            raise ProgrammeError('the workbook has no "Programme" sheet; use templates/programme_template.xlsx')
        return _sheet_rows(wb["Programme"]), (_sheet_rows(wb["Races"]) if "Races" in wb.sheetnames else [])
    with open(path, newline="", encoding="utf-8-sig") as f:
        return [{(k or "").strip().lower(): (v or "").strip() for k, v in r.items()} for r in csv.DictReader(f)], []


def load(conn, rid: int, path: str, replace: bool = False) -> dict:
    rows_, races = read_rows(path)
    out = load_rows(conn, rid, rows_, replace)
    added = 0
    for i, r in enumerate(races, start=2):
        try:
            d = date.fromisoformat(r.get("date", ""))
        except ValueError:
            raise ProgrammeError(f"Races row {i}: date must be YYYY-MM-DD") from None
        if not r.get("name"):
            continue
        if conn.execute("SELECT 1 FROM races WHERE runner_id=? AND name=? AND date=?",
                        (rid, r["name"], d.isoformat())).fetchone():
            continue
        conn.execute("INSERT INTO races (runner_id, name, date, distance_km, ascent_m, note) VALUES (?,?,?,?,?,?)",
                     (rid, r["name"][:120], d.isoformat(), float(r["distance_km"]) if r.get("distance_km") else None,
                      float(r["ascent_m"]) if r.get("ascent_m") else None, r.get("note") or None))
        added += 1
    conn.commit()
    if races:
        out["races_added"] = added
    return out


def load_csv(conn, rid: int, path: str, replace: bool = False) -> dict:
    return load(conn, rid, path, replace)


def load_rows(conn, rid: int, rows_: list[dict], replace: bool = False) -> dict:
    if not rows_:
        raise ProgrammeError("the file has no sessions")
    by_type: dict[str, list[dict]] = {}
    for i, r in enumerate(rows_, start=2):
        try:
            d = date.fromisoformat(r.get("date", ""))
        except ValueError:
            raise ProgrammeError(f"line {i}: date must be YYYY-MM-DD") from None
        t = r.get("type", "run") or "run"
        if t not in ("run", "strength"):
            raise ProgrammeError(f"line {i}: type must be run or strength")
        if not r.get("description"):
            raise ProgrammeError(f"line {i}: description is empty")
        target: dict = {}
        if r.get("km"):
            target["km"] = float(r["km"])
        target["kind"] = (r.get("kind") or ("strength" if t == "strength" else "easy")).lower()
        if t == "run" and target["kind"] not in RUN_KINDS:
            raise ProgrammeError(f"line {i}: kind must be one of {', '.join(RUN_KINDS)}")
        if r.get("climb_m"):
            target["climb_m"] = float(r["climb_m"])
        if r.get("hr_low") and r.get("hr_high"):
            target["hr_range"] = [int(r["hr_low"]), int(r["hr_high"])]
        elif r.get("hr_high"):
            target["hr_max"] = int(r["hr_high"])
        if r.get("steps"):
            # Structured workout for the watch, in Intervals.icu's workout text; "|" starts a new line.
            target["steps"] = "\n".join(x.strip() for x in r["steps"].split("|") if x.strip())
        if r.get("week"):
            target["week"] = int(float(r["week"]))
        if r.get("phase"):
            target["phase"] = r["phase"][:40]
        if r.get("exercises"):
            target["exercises"] = [e.strip() for e in r["exercises"].split(";") if e.strip()]
        by_type.setdefault(t, []).append({"date": d.isoformat(), "description": r["description"][:300],
                                          "target": target})
    out = {}
    for t, sessions in by_type.items():
        if replace:
            conn.execute("UPDATE plans SET active=0 WHERE runner_id=? AND type=? AND active=1", (rid, t))
        elif conn.execute("SELECT 1 FROM plans WHERE runner_id=? AND type=? AND active=1 AND end_date>=date('now')",
                          (rid, t)).fetchone():
            raise ProgrammeError(f"a {t} programme is already active; use --replace to end it first")
        sessions.sort(key=lambda s: s["date"])
        pid = conn.execute("INSERT INTO plans (runner_id, type, start_date, end_date) VALUES (?,?,?,?)",
                           (rid, t, sessions[0]["date"], sessions[-1]["date"])).lastrowid
        conn.executemany("INSERT INTO planned_workouts (plan_id, runner_id, date, description, target)"
                         " VALUES (?,?,?,?,?)",
                         [(pid, rid, s["date"], s["description"], json.dumps(s["target"])) for s in sessions])
        out[t] = {"plan_id": pid, "sessions": len(sessions), "from": sessions[0]["date"], "to": sessions[-1]["date"]}
    conn.commit()
    return out
