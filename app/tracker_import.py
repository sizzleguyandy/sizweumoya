"""Import Coach Andy's race tracker workbook (Race Info, Weekly Plan, Run Log, Strength tabs).

The Weekly Tracker tab is not imported: the system works it out from the runs.
Runs from the Run Log are stored with source='sheet'. When the runner later sends the
.fit file for the same run, the file replaces the spreadsheet row (see ingest.store_run).
"""
import hashlib
import json
import re
from datetime import date, datetime, timedelta

import openpyxl

from . import plans

# Tables holding a runner's data, cleared by reset_runner (the runner row itself stays).
RUNNER_TABLES = ("strength_sets", "strength_sessions", "run_splits", "runs", "weekly_summary", "fitness_checks",
                 "messages", "tips_given", "coach_notes", "runner_summary", "approvals",
                 "planned_workouts", "plans", "coach_alerts", "races")

DAY_OFFSET = {"Tue": 1, "Thu": 3, "Fri": 4, "Sun": 6}


class ImportError_(ValueError):
    pass


def reset_runner(conn, rid: int) -> dict:
    """Delete every activity, note, plan and message for one runner. Identity is kept."""
    counts = {}
    for t in RUNNER_TABLES:
        counts[t] = conn.execute(f"DELETE FROM {t} WHERE runner_id=?", (rid,)).rowcount
    conn.execute("UPDATE runners SET max_hr=NULL, resting_hr=NULL, hr_zones=NULL, zones_locked=0,"
                 " easy_hr_cap=NULL, weight_kg=NULL, goal=NULL WHERE runner_id=?", (rid,))
    conn.commit()
    return {k: v for k, v in counts.items() if v}


def has_data(conn, rid: int) -> bool:
    return any(conn.execute(f"SELECT 1 FROM {t} WHERE runner_id=? LIMIT 1", (rid,)).fetchone()
               for t in ("runs", "strength_sessions", "plans", "coach_notes"))


def _num(text, pattern=r"(\d+(?:\.\d+)?)"):
    m = re.search(pattern, str(text or ""))
    return float(m.group(1)) if m else None


def _day(v) -> date | None:
    return v.date() if isinstance(v, datetime) else v if isinstance(v, date) else None


def _key(*parts) -> str:
    return "sheet:" + hashlib.sha256("|".join(map(str, parts)).encode()).hexdigest()


# ---- Race Info --------------------------------------------------------------

def read_race_info(ws) -> dict:
    kv, zones = {}, []
    title = None
    for a, b in ws.iter_rows(max_col=2, values_only=True):
        if a is None:
            continue
        a = str(a).strip()
        if title is None:
            title = a
        if re.match(r"^[1-5]\s+\S", a) and b:
            lo_hi = re.findall(r"\d+", str(b))[:2]
            if len(lo_hi) == 2:
                zones.append((int(lo_hi[0]), int(lo_hi[1])))
        elif b is not None:
            kv[a.lower()] = str(b)
    if len(zones) != 5:
        raise ImportError_("Race Info: expected five heart rate zone rows")
    race_date = datetime.strptime(kv["date"].strip(), "%A %d %B %Y").date()
    name = re.split(r"\s+[—-]\s+", title or "Race")[0].title()
    easy_cap = _num(kv.get("easy hr guardrail"), r"HR\s*(\d+)")
    note = "; ".join(f"{k.capitalize()}: {kv[k]}" for k in ("start time", "location", "surface", "route")
                     if k in kv)
    return {
        "race": {"name": name, "date": race_date.isoformat(), "distance_km": _num(kv.get("distance")),
                 "ascent_m": _num(kv.get("total ascent")), "note": note or None},
        "zones": [hi for _, hi in zones[:4]],
        "max_hr": zones[4][1],
        "easy_hr_cap": int(easy_cap) if easy_cap else None,
        "weight_kg": _num(kv.get("weight")),
        "hm_range": [int(x) for x in re.findall(r"\d{3}", kv.get("hm effort", ""))[:2]] or None,
        "hill_range": [int(x) for x in re.findall(r"\d{3}", kv.get("hill reps", ""))[:2]] or None,
    }


# ---- Weekly Plan ------------------------------------------------------------

def _quality_kind(desc: str) -> str:
    d = desc.upper()
    if "TIME TRIAL" in d:
        return "benchmark"
    if d.startswith("EASY") or "SHAKEOUT" in d:
        return "easy"
    return "workout"


def read_plan(ws, info: dict) -> tuple[list[dict], list[str]]:
    rows = list(ws.iter_rows(values_only=True))
    head = [str(h or "") for h in rows[0]]
    col = {h.split("\n")[0].split(" —")[0].strip() + ("" if "\n" not in h else " " + h.split("\n")[1]): i
           for i, h in enumerate(head)}

    def c(name):
        if name in col:
            return col[name]
        for k, i in col.items():
            if k.startswith(name):
                return i
        raise ImportError_(f"Weekly Plan: no column {name!r}")

    i_start, i_phase = c("Week Start"), c("Phase")
    i_tue, i_thu_d, i_thu, i_fri = c("Tue Easy"), c("Thu"), c("Thu (km)"), c("Fri Easy")
    i_sun_d, i_sun, i_climb = c("Sun"), c("Sun (km)"), c("Sun climb")
    done = {d: c(f"{d} Done") for d in DAY_OFFSET}
    workouts, notes = [], []
    for r in rows[1:]:
        start = _day(r[i_start])
        if start is None:
            if r[0] and isinstance(r[0], str) and len(r[0]) > 40:
                notes.append(r[0])
            continue
        week, phase = int(r[0]), r[i_phase]

        def add(day, desc, km, kind, **extra):
            if not isinstance(km, (int, float)):
                return
            t = {"week": week, "phase": phase, "day": day, "km": km, "kind": kind, **extra}
            if kind == "workout":
                rng = info["hm_range"] if "HM EFFORT" in desc.upper() or "HILL" not in desc.upper() \
                    else info["hill_range"]
                if rng:
                    t["hr_range"] = rng
            if "HM effort" in desc and kind in ("long",) and info["hm_range"]:
                t["hr_range"] = info["hm_range"]
            workouts.append({"date": (start + timedelta(days=DAY_OFFSET[day])).isoformat(),
                             "description": desc, "target": t, "ticked": bool(r[done[day]])})

        add("Tue", "Easy run", r[i_tue], "easy")
        add("Thu", str(r[i_thu_d]), r[i_thu], _quality_kind(str(r[i_thu_d])))
        add("Fri", "Easy run", r[i_fri], "easy")
        sun = str(r[i_sun_d])
        is_race = sun.upper().startswith("RACE")
        add("Sun", sun, r[i_sun], "race" if is_race else "long",
            **({"climb_m": r[i_climb]} if isinstance(r[i_climb], (int, float)) else {}))
    for w in workouts:
        if w["target"]["kind"] == "race":
            w["target"].pop("hr_range", None)
    return workouts, notes


# ---- Run Log ----------------------------------------------------------------

SESSION_TYPE = {"easy": "easy", "long": "long", "quality": "workout", "sick recovery": "easy"}


def read_run_log(ws) -> tuple[list[dict], list[dict], list[dict]]:
    """Returns (runs, sick days, gym days)."""
    runs, sick, gym = [], [], []
    for r in list(ws.iter_rows(values_only=True))[1:]:
        d = _day(r[0])
        if d is None:
            continue
        session, km, mins, hr, elev, note = (str(r[2] or "").strip(), r[3], r[4], r[6], r[7], r[8])
        low = session.lower()
        if low == "sick":
            sick.append({"date": d.isoformat(), "note": note})
            continue
        if low == "gym":
            gym.append({"date": d.isoformat(), "note": note})
            continue
        if not isinstance(km, (int, float)) or not km:
            continue
        run_type = SESSION_TYPE.get(low, "easy")
        if run_type == "workout" and note and re.search(r"all-out|time trial", note, re.I):
            run_type = "benchmark"
        dur = int(round(mins * 60)) if isinstance(mins, (int, float)) else None
        runs.append({"date": d.isoformat(), "distance_km": round(km, 2), "duration_s": dur,
                     "avg_pace_s_km": round(dur / km, 1) if dur else None,
                     "avg_hr": hr if isinstance(hr, (int, float)) else None,
                     "elevation_m": elev if isinstance(elev, (int, float)) else None,
                     "run_type": run_type, "note": (note or None),
                     "treadmill": int(bool(note and "treadmill" in note.lower()))})
    return runs, sick, gym


# ---- Strength ---------------------------------------------------------------

def read_strength(ws) -> dict[str, list[dict]]:
    sessions: dict[str, list[dict]] = {}
    started = False
    for r in ws.iter_rows(values_only=True):
        if not started:
            started = r[0] == "Date" and r[2] == "Exercise"
            continue
        d = _day(r[0])
        if d is None or not r[2]:
            continue
        sessions.setdefault(d.isoformat(), []).append({
            "exercise": str(r[2]).strip().lower(), "set_no": int(r[3] or 0),
            "reps": int(r[4]) if isinstance(r[4], (int, float)) else None,
            "weight_kg": float(r[5]) if isinstance(r[5], (int, float)) else None,
            "note": r[8] if len(r) > 8 else None})
    return sessions


# ---- import -----------------------------------------------------------------

def import_workbook(conn, runner, path: str, today: date | None = None) -> dict:
    today = today or date.today()
    wb = openpyxl.load_workbook(path, data_only=True)
    for name in ("Race Info", "Weekly Plan", "Run Log"):
        if name not in wb.sheetnames:
            raise ImportError_(f"the workbook has no {name!r} tab")
    rid = runner["runner_id"]
    info = read_race_info(wb["Race Info"])
    workouts, plan_notes = read_plan(wb["Weekly Plan"], info)
    runs, sick, gym = read_run_log(wb["Run Log"])
    lifts = read_strength(wb["Strength"]) if "Strength" in wb.sheetnames else {}

    race = info["race"]
    conn.execute("UPDATE runners SET max_hr=?, hr_zones=?, zones_locked=1, easy_hr_cap=?, weight_kg=?, goal=?"
                 " WHERE runner_id=?",
                 (info["max_hr"], json.dumps(info["zones"]), info["easy_hr_cap"], info["weight_kg"],
                  f"{race['name']} {race['distance_km']} km, {race['date']}", rid))
    conn.execute("INSERT INTO races (runner_id, name, date, distance_km, ascent_m, note) VALUES (?,?,?,?,?,?)",
                 (rid, race["name"], race["date"], race["distance_km"], race["ascent_m"], race["note"]))

    pid = conn.execute("INSERT INTO plans (runner_id, type, start_date, end_date) VALUES (?,?,?,?)",
                       (rid, "run", workouts[0]["date"], workouts[-1]["date"])).lastrowid
    for w in workouts:
        conn.execute("INSERT INTO planned_workouts (plan_id, runner_id, date, description, target)"
                     " VALUES (?,?,?,?,?)", (pid, rid, w["date"], w["description"], json.dumps(w["target"])))
    if plan_notes:
        conn.execute("INSERT INTO coach_notes (runner_id, source, category, note) VALUES (?,?,?,?)",
                     (rid, "plan", "preference", "Plan notes from Coach Andy: " + " | ".join(plan_notes)[:2900]))
    conn.commit()

    runner = conn.execute("SELECT * FROM runners WHERE runner_id=?", (rid,)).fetchone()
    for r in runs:
        row = {"runner_id": rid, "activity_key": _key(rid, r["date"], r["distance_km"], r["duration_s"]),
               "source": "sheet", "name": "Treadmill run" if r["treadmill"] else "Run", "is_easy_valid": 0, **r}
        cols = ", ".join(row)
        conn.execute(f"INSERT OR IGNORE INTO runs ({cols}) VALUES ({', '.join('?' * len(row))})",
                     list(row.values()))
        conn.commit()
        plans.match_run(conn, runner, row)

    for d, sets in lifts.items():
        sid = conn.execute(
            "INSERT INTO strength_sessions (runner_id, activity_key, source, date, title, parse_ok)"
            " VALUES (?,?,?,?,?,1)", (rid, _key(rid, "strength", d), "sheet", d, "Gym session")).lastrowid
        conn.executemany("INSERT INTO strength_sets (session_id, runner_id, exercise, set_no, reps, weight_kg)"
                         " VALUES (?,?,?,?,?,?)",
                         [(sid, rid, s["exercise"], s["set_no"], s["reps"], s["weight_kg"]) for s in sets])
    for g in gym:
        if g["date"] not in lifts:
            conn.execute("INSERT INTO strength_sessions (runner_id, activity_key, source, date, title, note)"
                         " VALUES (?,?,?,?,?,?)",
                         (rid, _key(rid, "strength", g["date"]), "sheet", g["date"], "Gym session", g["note"]))
    if sick:
        days = ", ".join(s["date"] for s in sick)
        conn.execute("INSERT INTO coach_notes (runner_id, source, category, note, status) VALUES (?,?,?,?,?)",
                     (rid, "sheet", "health", f"Illness: missed sessions on {days}. "
                      + " ".join(dict.fromkeys(s["note"] for s in sick if s["note"])), "resolved"))

    # Past sessions with no run: ticked in the sheet counts as done, otherwise missed.
    for w in workouts:
        if w["date"] >= (today - timedelta(days=plans.MATCH_DAYS)).isoformat():
            continue
        status = "done" if w["ticked"] else "missed"
        conn.execute("UPDATE planned_workouts SET status=? WHERE plan_id=? AND date=? AND status='planned'",
                     (status, pid, w["date"]))
    conn.commit()
    done = conn.execute("SELECT status, COUNT(*) FROM planned_workouts WHERE plan_id=? GROUP BY status",
                        (pid,)).fetchall()
    return {"race": race, "zones": info["zones"], "max_hr": info["max_hr"], "easy_hr_cap": info["easy_hr_cap"],
            "planned_sessions": len(workouts), "plan_status": {s: n for s, n in done},
            "runs": len(runs), "strength_sessions": len(lifts) + sum(g["date"] not in lifts for g in gym),
            "sick_days": len(sick)}
