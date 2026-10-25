"""File uploads: .fit runs and strength sessions, zips, typed workouts, safety and migration."""
import io
import json
import sqlite3
import zipfile
from dataclasses import replace
from datetime import date, datetime, timedelta, timezone

import pytest

from app import config, ingest, queries, upload, workout_text
from app.admin import add_runner
from app.db import connect, init_db
from tests.fitgen import run_fit, strength_fit

UTC = timezone.utc


@pytest.fixture
def env(tmp_path, monkeypatch):
    home = tmp_path / "hermes"
    monkeypatch.setattr(upload, "settings", replace(config.settings, hermes_home=str(home)))
    conn = connect(":memory:")
    init_db(conn)
    rid, profile, _ = add_runner(conn, "Tom", "111", 185, "Sub-2 half")
    folder = home / "profiles" / profile
    folder.mkdir(parents=True)
    runner = conn.execute("SELECT * FROM runners WHERE runner_id=?", (rid,)).fetchone()
    yield conn, runner, folder, home
    conn.close()


def _save(folder, name, data):
    p = folder / name
    p.write_bytes(data)
    return str(p)


def _start(days_ago=1, hour=4):
    d = datetime.now(UTC).replace(hour=hour, minute=30, second=0, microsecond=0)
    return d - timedelta(days=days_ago)


def test_run_file_is_stored_with_stats(env):
    conn, runner, folder, _ = env
    data = run_fit(_start(), seconds=2400, speed_mps=3.0, hr=140)
    res = upload.process_path(conn, runner, _save(folder, "run.fit", data), note="felt good")
    act = res["activity"]
    assert act["status"] == "stored" and act["kind"] == "run"
    assert act["distance_km"] == 7.2 and act["duration_s"] == 2400 and act["short"] is False
    row = conn.execute("SELECT * FROM runs").fetchone()
    assert row["note"] == "felt good" and row["avg_pace_s_km"] == pytest.approx(333.3, abs=0.1)
    assert row["z2_s"] > 2000  # 140 bpm sits in zone 2 for max HR 185


def test_same_file_twice_is_stored_once(env):
    conn, runner, folder, _ = env
    path = _save(folder, "run.fit", run_fit(_start(), seconds=1500))
    upload.process_path(conn, runner, path)
    again = upload.process_path(conn, runner, path)
    assert again["activity"]["status"] == "duplicate" and again["duplicates"] == 1
    assert conn.execute("SELECT COUNT(*) FROM runs").fetchone()[0] == 1


def test_same_run_from_another_file_is_stored_once(env):
    """The watch's export and Intervals.icu's copy of one run are different files."""
    conn, runner, folder, _ = env
    upload.process_path(conn, runner, _save(folder, "a.fit", run_fit(_start(), seconds=2348, hr=158)))
    twin = upload.process_path(conn, runner, _save(folder, "b.fit", run_fit(_start(), seconds=2349, hr=157)))
    assert twin["activity"]["status"] == "duplicate" and twin["activity"]["same_as"].startswith("fit:")
    other = upload.process_path(conn, runner, _save(folder, "c.fit", run_fit(_start(hour=16), seconds=1200)))
    assert other["activity"]["status"] == "stored"      # a second, different run that day is kept
    assert conn.execute("SELECT COUNT(*) FROM runs").fetchone()[0] == 2


def test_dedupe_keeps_the_run_the_programme_points_to(env):
    conn, runner, _, _ = env
    rid = runner["runner_id"]
    d = _start().date().isoformat()
    for key, secs in (("fit:a", 2348), ("fit:b", 2349), ("fit:c", 1200)):
        conn.execute("INSERT INTO runs (runner_id, activity_key, date, distance_km, duration_s) VALUES (?,?,?,?,?)",
                     (rid, key, d, 6.18 if secs > 2000 else 3.0, secs))
    pid = conn.execute("INSERT INTO plans (runner_id, type, start_date, active) VALUES (?, 'run', ?, 1)",
                       (rid, d)).lastrowid
    conn.execute("INSERT INTO planned_workouts (plan_id, runner_id, date, description, status,"
                 " completed_activity_key) VALUES (?,?,?,?, 'done', 'fit:b')", (pid, rid, d, "Easy"))
    pairs = ingest.duplicate_runs(conn, rid)
    assert [(k["activity_key"], x["activity_key"]) for k, x in pairs] == [("fit:b", "fit:a")]
    assert ingest.remove_duplicate_runs(conn, rid) == 1
    assert {r[0] for r in conn.execute("SELECT activity_key FROM runs")} == {"fit:b", "fit:c"}


def test_short_run_flagged(env):
    conn, runner, folder, _ = env
    res = upload.process_path(conn, runner, _save(folder, "s.fit", run_fit(_start(), seconds=600, speed_mps=3)))
    assert res["activity"]["short"] is True  # 1.8 km in 10 min


def test_treadmill_never_counts_for_fitness_check(env):
    conn, runner, folder, _ = env
    upload.process_path(conn, runner, _save(folder, "t.fit",
                        run_fit(_start(), seconds=2400, hr=135, sub_sport="treadmill")))
    row = conn.execute("SELECT treadmill, is_easy_valid, name FROM runs").fetchone()
    assert row["treadmill"] == 1 and row["is_easy_valid"] == 0 and row["name"] == "Treadmill run"


def test_caption_classifies_race(env):
    conn, runner, folder, _ = env
    upload.process_path(conn, runner, _save(folder, "r.fit", run_fit(_start(), seconds=1500, hr=175)),
                        note="parkrun PB!")
    assert conn.execute("SELECT run_type FROM runs").fetchone()[0] == "race"


def test_strength_file_stores_sets(env):
    conn, runner, folder, _ = env
    data = strength_fit(_start(hour=15), [("squat", 8, 60), ("squat", 8, 62.5), ("deadlift", 5, 100)])
    res = upload.process_path(conn, runner, _save(folder, "gym.fit", data))
    assert res["activity"]["kind"] == "strength" and res["activity"]["sets"] == 3
    last = queries.last_strength_session(conn, runner["runner_id"])
    assert [(s["exercise"], s["reps"], s["weight_kg"]) for s in last["sets"]] == [
        ("squat", 8, 60.0), ("squat", 8, 62.5), ("deadlift", 5, 100.0)]


def test_zip_history_import_and_baseline(env):
    conn, runner, folder, _ = env
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for i in range(8):
            z.writestr(f"a{i}.fit", run_fit(_start(days_ago=2 + i * 5), seconds=2700, hr=140))
        z.writestr(f"gym.fit", strength_fit(_start(days_ago=3, hour=16), [("squat", 8, 60)]))
        z.writestr("readme.txt", "ignored")
    res = upload.process_path(conn, runner, _save(folder, "export.zip", buf.getvalue()))
    assert res["history_import"] and res["stored_runs"] == 8 and res["stored_strength"] == 1
    base = queries.baseline(conn, runner["runner_id"])
    assert base["runs"] == 8 and base["enough_data"] and base["longest_run_km"] == 8.1
    assert base["easy_pct"] >= 90 and base["easy_hr"] == 140 and base["strength_sessions"] == 1


def test_path_outside_own_folder_is_refused(env, tmp_path):
    conn, runner, folder, home = env
    outside = tmp_path / "secret.fit"
    outside.write_bytes(run_fit(_start(), seconds=60))
    other = home / "profiles" / "runner-99"
    other.mkdir(parents=True)
    other_file = other / "theirs.fit"
    other_file.write_bytes(run_fit(_start(), seconds=60))
    for bad in (str(outside), str(other_file), str(folder / ".." / "runner-99" / "theirs.fit"),
                str(folder / "missing.fit")):
        with pytest.raises(upload.UploadRejected):
            upload.process_path(conn, runner, bad)
    assert conn.execute("SELECT COUNT(*) FROM runs").fetchone()[0] == 0


def test_wrong_type_and_garbage(env):
    conn, runner, folder, _ = env
    with pytest.raises(upload.UploadRejected):
        upload.process_path(conn, runner, _save(folder, "notes.txt", b"hello"))
    res = upload.process_path(conn, runner, _save(folder, "bad.fit", b"not a fit file"))
    assert res["activity"]["status"] == "unreadable"


def test_uploads_refused_when_not_configured(env, monkeypatch):
    conn, runner, folder, _ = env
    monkeypatch.setattr(upload, "settings", replace(config.settings, hermes_home=""))
    with pytest.raises(upload.UploadRejected, match="set up"):
        upload.process_path(conn, runner, _save(folder, "r.fit", run_fit(_start(), seconds=60)))


def test_max_hr_raised_unless_coach_locked(env):
    conn, runner, folder, _ = env
    upload.process_path(conn, runner, _save(folder, "hard.fit",
                        run_fit(_start(), seconds=900, hr=lambda i: 150 + i // 20)))
    r = conn.execute("SELECT max_hr, hr_zones FROM runners").fetchone()
    assert r["max_hr"] == 195 and json.loads(r["hr_zones"])[0] == round(195 * 0.68)
    conn.execute("UPDATE runners SET zones_locked=1, hr_zones='[100,110,120,130]'")
    runner = conn.execute("SELECT * FROM runners").fetchone()
    upload.process_path(conn, runner, _save(folder, "harder.fit",
                        run_fit(_start(days_ago=2), seconds=900, hr=lambda i: 160 + i // 20)))
    r = conn.execute("SELECT max_hr, hr_zones FROM runners").fetchone()
    assert r["max_hr"] == 205 and r["hr_zones"] == "[100,110,120,130]"


def test_typed_workout():
    sets, unparsed = workout_text.parse("Squat 3x8 60kg\nRDL 3 x 10 @ 50 kg; Plank 3x45s\nfelt strong")
    assert [(s.exercise, s.reps, s.weight_kg, s.duration_s) for s in sets[:1]] == [("squat", 8, 60.0, None)]
    assert len(sets) == 9 and sets[3].exercise == "rdl" and sets[3].weight_kg == 50.0
    assert sets[-1].duration_s == 45 and unparsed == ["felt strong"]


def test_typed_workout_stored_once(env):
    conn, runner, _, _ = env
    today = date.today().isoformat()
    first = ingest.store_strength_text(conn, runner, "Squat 3x8 60kg", today)
    assert first["sets"] == 3 and not first["duplicate"]
    assert ingest.store_strength_text(conn, runner, "Squat 3x8 60kg", today)["duplicate"]


def test_migration_from_strava_schema(tmp_path):
    path = tmp_path / "old.db"
    old = sqlite3.connect(path)
    old.executescript("""
        CREATE TABLE runners (runner_id INTEGER PRIMARY KEY, name TEXT NOT NULL, telegram_chat_id TEXT,
            hermes_profile TEXT UNIQUE NOT NULL, mcp_token_hash TEXT NOT NULL, hermes_api_key TEXT,
            tier TEXT DEFAULT 'free', strava_athlete_id INTEGER, max_hr INTEGER, resting_hr INTEGER,
            hr_zones TEXT, goal TEXT, units TEXT DEFAULT 'km', joined TEXT DEFAULT (date('now')),
            active INTEGER DEFAULT 1);
        CREATE TABLE runs (id INTEGER PRIMARY KEY, runner_id INTEGER NOT NULL, strava_id INTEGER NOT NULL UNIQUE,
            date TEXT NOT NULL, name TEXT, distance_km REAL, duration_s INTEGER, avg_pace_s_km REAL,
            avg_hr REAL, max_hr REAL, z1_s INTEGER, z2_s INTEGER, z3_s INTEGER, z4_s INTEGER, z5_s INTEGER,
            elevation_m REAL, temp_c REAL, hr_drift_pct REAL, run_type TEXT, is_easy_valid INTEGER DEFAULT 0);
        CREATE TABLE strength_sessions (id INTEGER PRIMARY KEY, runner_id INTEGER NOT NULL,
            strava_id INTEGER NOT NULL UNIQUE, date TEXT NOT NULL, title TEXT, duration_s INTEGER,
            avg_hr REAL, raw_description TEXT, parse_ok INTEGER DEFAULT 0);
        CREATE TABLE strength_sets (id INTEGER PRIMARY KEY, session_id INTEGER NOT NULL
            REFERENCES strength_sessions(id) ON DELETE CASCADE, runner_id INTEGER NOT NULL, exercise TEXT,
            set_no INTEGER, reps INTEGER, weight_kg REAL, duration_s INTEGER, distance_m REAL);
        CREATE TABLE planned_workouts (id INTEGER PRIMARY KEY, plan_id INTEGER, runner_id INTEGER,
            date TEXT, description TEXT, target TEXT, completed_strava_id INTEGER, status TEXT);
        INSERT INTO runners (runner_id, name, hermes_profile, mcp_token_hash) VALUES (1, 'Tom', 'runner-1', 'x');
        INSERT INTO runs (runner_id, strava_id, date, distance_km) VALUES (1, 555, '2026-10-01', 8.0);
        INSERT INTO strength_sessions (id, runner_id, strava_id, date) VALUES (7, 1, 777, '2026-10-02');
        INSERT INTO strength_sets (session_id, runner_id, exercise) VALUES (7, 1, 'squat');
    """)
    old.commit()
    old.close()
    conn = connect(str(path))
    init_db(conn)
    assert tuple(conn.execute("SELECT activity_key, source FROM runs").fetchone()) == ("strava:555", "strava")
    assert conn.execute("SELECT activity_key FROM strength_sessions").fetchone()[0] == "strava:777"
    assert conn.execute("SELECT zones_locked FROM runners").fetchone()[0] == 0
    assert conn.execute("PRAGMA foreign_key_check").fetchall() == []
    # FK from strength_sets still points at the rebuilt table
    conn.execute("DELETE FROM strength_sessions WHERE id=7")
    assert conn.execute("SELECT COUNT(*) FROM strength_sets").fetchone()[0] == 0
    init_db(conn)  # running again is harmless
    conn.close()



def _tracker(path):
    """A small workbook in the same layout as Coach Andy's race tracker."""
    import openpyxl
    wb = openpyxl.Workbook()
    ri = wb.active
    ri.title = "Race Info"
    for row in [("TEST HALF — 21.1 km", None), ("Date", "Sunday 15 November 2026"), ("Start time", "05:45"),
                ("Distance", "21.1 km"), ("Total ascent", "445 m"), ("Weight", "97 kg"),
                ("Easy HR guardrail", "Keep easy runs under HR 138."), ("HM effort", "HR 152 - 163"),
                ("Hill reps", "HR 160 - 175 on the climb"), ("1  Recovery", "109 - 135"), ("2  Easy", "136 - 147"),
                ("3  Aerobic", "148 - 156"), ("4  Threshold", "157 - 166"), ("5  Anaerobic", "167 - 184")]:
        ri.append(row)
    wp = wb.create_sheet("Weekly Plan")
    wp.append(["Wk", "Week Start", "Week End", "Phase", "Runs", "Tue\nEasy (km)",
               "Thu — Quality Session", "Thu\n(km)", "Fri\nEasy (km)", "Sun — Long Run", "Sun\n(km)",
               "Sun climb\ntarget (m)", "Total\n(km)", "Tue Done", "Thu Done", "Fri Done", "Sun Done"])
    wp.append([1, datetime(2026, 10, 12), datetime(2026, 10, 18), "Build", 4, 5, "5K TIME TRIAL — flat", 8, 5,
               "Hilly. HR <138", 12, 150, 30, True, True, False, False])
    wp.append([2, datetime(2026, 10, 19), datetime(2026, 10, 25), "Build", 4, 5,
               "HILLY HM REPS — 2 x 2km @ HM effort", 7, 5, "Hilly. Last 20min @ HM effort", 16, 250, 33,
               False, False, False, False])
    rl = wb.create_sheet("Run Log")
    rl.append(["Date", "Day", "Session", "Distance (km)", "Time (min)", "Pace", "Avg HR", "Elevation (m)", "Notes"])
    rl.append([datetime(2026, 10, 13), "Tue", "Easy", 5.0, 35.0, None, 135, 40, None])
    rl.append([datetime(2026, 10, 14), "Wed", "Sick", None, None, None, None, None, "cold"])
    rl.append([datetime(2026, 10, 17), "Sat", "Easy", 4.26, 34.8, None, 130, 0, "treadmill"])
    rl.append([datetime(2026, 10, 18), "Sun", "Long", 10.07, 68.6, None, 151, 155, "ran hot"])
    st = wb.create_sheet("Strength")
    st.append(["LIFT LOG"])
    st.append(["Date", "Day", "Exercise", "Set", "Reps", "Weight (kg)"])
    st.append([datetime(2026, 10, 15), "Thu", "Chest Press (Machine)", 1, 5, 35])
    st.append([datetime(2026, 10, 15), "Thu", "Chest Press (Machine)", 2, 7, 35])
    wb.save(path)


def test_tracker_import_and_fit_replaces_sheet_row(env, tmp_path):
    from app import plans, tracker_import
    conn, runner, folder, _ = env
    rid = runner["runner_id"]
    queries.add_note(conn, rid, "pattern", "old test note")
    assert tracker_import.has_data(conn, rid)
    tracker_import.reset_runner(conn, rid)
    assert not tracker_import.has_data(conn, rid)
    path = tmp_path / "t.xlsx"
    _tracker(path)
    res = tracker_import.import_workbook(conn, runner, str(path), today=date(2026, 10, 21))
    assert res["zones"] == [135, 147, 156, 166] and res["max_hr"] == 184 and res["easy_hr_cap"] == 138
    assert res["runs"] == 3 and res["sick_days"] == 1 and res["strength_sessions"] == 1
    r = conn.execute("SELECT zones_locked, weight_kg FROM runners").fetchone()
    assert tuple(r) == (1, 97.0)
    assert plans.races(conn, rid, date(2026, 10, 21))[0]["days_to_go"] == 25
    status = dict(conn.execute("SELECT date, status FROM planned_workouts").fetchall())
    # Thu ticked with no run -> done; Sat treadmill run counts for Friday; Sun matched.
    assert status["2026-10-15"] == "done" and status["2026-10-16"] == "done" and status["2026-10-18"] == "done"
    assert status["2026-10-22"] == "planned"
    long_run = [x for x in queries.recent_runs(conn, rid, 5) if x["date"] == "2026-10-18"][0]
    assert long_run["run_type"] == "long" and long_run["plan"]["over_hr_cap"] and "over_hr_cap" in long_run["tips"]
    thu = conn.execute("SELECT target FROM planned_workouts WHERE date='2026-10-22'").fetchone()[0]
    assert json.loads(thu)["hr_range"] == [152, 163]

    # The .fit for the Sunday run replaces the spreadsheet row and keeps the plan link.
    runner = conn.execute("SELECT * FROM runners").fetchone()
    start = datetime(2026, 10, 18, 9, 11, tzinfo=UTC)
    out = upload.process_path(conn, runner, _save(folder, "sun.fit",
                              run_fit(start, seconds=4116, speed_mps=2.447, hr=151)))
    act = out["activity"]
    assert act["replaced_sheet_row"] and act["plan"]["planned_date"] == "2026-10-18"
    rows_ = conn.execute("SELECT source, note FROM runs WHERE date='2026-10-18'").fetchall()
    assert [tuple(x) for x in rows_] == [("fit", "ran hot")]


def test_strength_review(env):
    from app import strength
    conn, runner, _, _ = env
    rid = runner["runner_id"]
    for i, d in enumerate(["2026-09-30", "2026-10-07", "2026-10-14", "2026-10-17", "2026-10-21"]):
        ingest.store_strength_text(conn, runner, "Chest press 3x8 35kg\nLat pulldown 3x5 30kg"
                                   + ("\nSquat 3x8 60kg" if i == 4 else ""), d)
    ingest.store_strength_text(conn, runner, "Lat pulldown 3x7 30kg", "2026-10-20")
    rev = strength.review(conn, rid, today=date(2026, 10, 21))
    vol = rev["last_7_days"]["by_muscle"]
    assert vol["chest"]["sets"] == 6 and vol["chest"]["days"] == 2 and vol["chest"]["volume"] == "below start range"
    assert vol["back"]["sets"] == 9 and vol["quads"]["sets"] == 3
    lifts = {x["exercise"]: x for x in rev["lifts"]}
    assert lifts["chest press"]["stalled"] and lifts["chest press"]["same_sessions_in_a_row"] == 5
    assert lifts["lat pulldown"]["trend"] == "reps down" and lifts["squat"]["trend"] == "first session"
    assert strength.muscles_for("Preacher Curl (Machine)") == ["biceps"]
    assert strength.muscles_for("Triceps Dip (Assisted)")[0] == "triceps"


def test_paces_and_zones_update_with_new_benchmark(env):
    from app import paces
    conn, runner, folder, _ = env
    rid = runner["runner_id"]
    conn.execute("UPDATE runners SET hr_zones='[132,144,152,162]', max_hr=184, easy_hr_cap=144, zones_locked=1")
    conn.commit()
    today = date.today()
    prof = paces.profile(conn, rid, today)
    assert prof["performance"]["available"] is False
    assert prof["hr_zones"][1] == {"zone": 2, "name": "Easy", "from": 133, "to": 144}

    def add(days_ago, km, secs, run_type, hr=140):
        d = (today - timedelta(days=days_ago)).isoformat()
        conn.execute("INSERT INTO runs (runner_id, activity_key, source, date, distance_km, duration_s,"
                     " avg_pace_s_km, avg_hr, run_type) VALUES (?,?,?,?,?,?,?,?,?)",
                     (rid, f"t:{days_ago}:{km}", "sheet", d, km, secs, secs / km, hr, run_type))
        conn.commit()

    add(30, 5.0, 30 * 60, "benchmark", 170)
    add(10, 6.0, 6 * 430, "easy", 138)
    add(5, 8.0, 8 * 425, "long", 140)
    prof = paces.profile(conn, rid, today)
    perf = prof["performance"]
    assert perf["available"] and 30.5 <= perf["vdot"] <= 31.5 and perf["race_estimates"]["5K"] == "30:00"
    assert prof["observed_easy"]["pace"] == "7:08/km"
    old_vdot = perf["vdot"]
    add(2, 5.0, 28 * 60, "benchmark", 172)   # a faster time trial moves everything
    perf = paces.profile(conn, rid, today)["performance"]
    assert perf["vdot"] > old_vdot and perf["race_estimates"]["5K"] == "28:00"
    assert perf["training_paces"]["threshold"]["from"] < "6:13/km"


def test_mode_follows_what_is_on_file(env):
    from app import plans
    conn, runner, _, _ = env
    rid = runner["runner_id"]
    today = date(2026, 10, 22)
    assert plans.mode(conn, rid, today) == {"focus": "running", "mode": "monitor", "running": "monitor",
                                            "strength": "monitor", "run_programme": None,
                                            "strength_programme": None}
    s0 = plans.sessions_around(conn, rid, today)
    assert s0["mode"] == "monitor" and s0["sessions"] == [] and s0["today_is"]["weekday"] == "Thursday"
    pid = conn.execute("INSERT INTO plans (runner_id, type, start_date, end_date) VALUES (?,?,?,?)",
                       (rid, "run", "2026-08-31", "2026-11-15")).lastrowid
    for d, desc, km in [("2026-10-20", "Easy run", 5), ("2026-10-22", "Hill reps", 7), ("2026-10-23", "Easy run", 5)]:
        conn.execute("INSERT INTO planned_workouts (plan_id, runner_id, date, description, target) VALUES (?,?,?,?,?)",
                     (pid, rid, d, desc, json.dumps({"kind": "easy", "km": km})))
    conn.commit()
    m = plans.mode(conn, rid, today)
    assert m["mode"] == "guided" and m["run_programme"]["weeks_left"] == 4
    s = plans.sessions_around(conn, rid, today)
    assert s["today"][0]["description"] == "Hill reps" and s["last"][0]["date"] == "2026-10-20"
    assert [w["date"] for w in plans.tomorrow_session(conn, rid, today)] == ["2026-10-23"]
    # After the race the programme has ended: back to option 1.
    assert plans.mode(conn, rid, date(2026, 11, 16))["mode"] == "monitor"


def test_week_summary_zones_and_programme(env):
    conn, runner, folder, _ = env
    rid = runner["runner_id"]
    conn.execute("UPDATE runners SET hr_zones='[132,144,152,162]', max_hr=184, easy_hr_cap=144")
    conn.commit()
    start = queries.week_start(date.today())
    when = datetime.combine(start, datetime.min.time(), UTC).replace(hour=5)
    runner = conn.execute("SELECT * FROM runners").fetchone()
    upload.process_path(conn, runner, _save(folder, "a.fit", run_fit(when, seconds=1800, hr=140)))
    week = queries.week_summary(conn, rid, start)
    assert week["coaching_mode"] == "monitor" and "vs_programme" not in week
    easy = week["time_in_zones"][1]
    assert easy["name"] == "Easy" and easy["minutes"] == 30 and easy["pct"] == 100
    assert week["total_time"] == "30:00" and week["easy_runs_over_hr_cap"] == (0, 1)

    pid = conn.execute("INSERT INTO plans (runner_id, type, start_date, end_date) VALUES (?,?,?,?)",
                       (rid, "run", start.isoformat(), (start + timedelta(days=60)).isoformat())).lastrowid
    for off, kind, km in [(1, "easy", 5), (6, "long", 12)]:
        conn.execute("INSERT INTO planned_workouts (plan_id, runner_id, date, description, target, status)"
                     " VALUES (?,?,?,?,?,?)", (pid, rid, (start + timedelta(days=off)).isoformat(), kind,
                                               json.dumps({"kind": kind, "km": km, "week": 3}),
                                               "missed" if off == 1 else "planned"))
    conn.commit()
    vs = queries.week_summary(conn, rid, start)["vs_programme"]
    assert vs["planned_km"] == 17 and vs["actual_km"] == 5.4 and vs["sessions_planned"] == 2
    assert len(vs["missed"]) == 1 and vs["long_run"]["planned_km"] == 12


def test_week_counts_sessions_of_a_programme_replaced_midweek():
    from app import plans
    c = connect(":memory:")
    init_db(c)
    rid, _, _ = add_runner(c, "Andy", "111", 186, None)
    start = date(2026, 10, 19)

    def plan(active, first, sessions):
        pid = c.execute("INSERT INTO plans (runner_id, type, start_date, end_date, active) VALUES (?,?,?,?,?)",
                        (rid, "run", first, "2026-11-15", active)).lastrowid
        for d, km, st in sessions:
            c.execute("INSERT INTO planned_workouts (plan_id, runner_id, date, description, target, status)"
                      " VALUES (?,?,?,?,?,?)", (pid, rid, d, "run", json.dumps({"kind": "easy", "km": km}), st))
    plan(0, "2026-10-01", [("2026-10-20", 5, "done"), ("2026-10-22", 6, "done"),
                           ("2026-10-24", 9, "planned"), ("2026-10-25", 13, "planned")])
    plan(1, "2026-10-24", [("2026-10-24", 5, "done"), ("2026-10-25", 13, "done")])
    vs = plans.week_vs_programme(c, rid, start, start + timedelta(days=7), {"total_km": 27.2})
    assert vs["planned_km"] == 29 and vs["sessions_planned"] == 4 and vs["done"] == 4


def test_month_summary(env):
    conn, runner, folder, _ = env
    rid = runner["runner_id"]
    for d, km in [("2026-09-10", 5.0), ("2026-10-02", 6.0), ("2026-10-20", 10.0)]:
        conn.execute("INSERT INTO runs (runner_id, activity_key, date, distance_km, duration_s, run_type)"
                     " VALUES (?,?,?,?,?,?)", (rid, f"k{d}", d, km, int(km * 420), "easy"))
    conn.commit()
    m = queries.month_summary(conn, rid, 2026, 10)
    assert m["month"] == "October 2026" and m["total_km"] == 16.0 and m["previous_month_km"] == 5.0
    assert m["longest_run_km"] == 10.0 and sum(w["km"] for w in m["weeks"]) == 16.0
    assert m["progress_status"] == "grey" and m["coaching_mode"] == "monitor"
    assert queries.month_summary(conn, rid, 2026, 12)["month"] == "December 2026"


def test_purge_strava_leaves_everything_else(tmp_path, monkeypatch):
    import subprocess, sys
    dbp = tmp_path / "r.db"
    c = connect(str(dbp))
    init_db(c)
    rid, _, _ = add_runner(c, "Tom", "1", 185, None)
    for key, src in (("strava:1", "strava"), ("fit:abc", "fit")):
        c.execute("INSERT INTO runs (runner_id, activity_key, source, date, distance_km) VALUES (?,?,?,?,?)",
                  (rid, key, src, "2026-10-01", 5))
    c.commit()
    c.close()
    env = {**__import__("os").environ, "RAC_DB_PATH": str(dbp), "RAC_RECORDS_DIR": str(tmp_path / "rec")}
    out = subprocess.run([sys.executable, "-m", "app.admin", "purge-strava"], env=env, capture_output=True, text=True)
    assert out.returncode == 0, out.stderr
    c = connect(str(dbp))
    assert [r[0] for r in c.execute("SELECT activity_key FROM runs")] == ["fit:abc"]


HEVY_SHARE = """Afternoon workout 💪
Friday, Oct 23, 2026 at 12:36pm

Sumo Deadlift
Set 1: 60 kg x 5
Set 2: 90 kg x 5
Set 3: 90 kg x 5
Set 4: 90 kg x 5

Shoulder Press (Machine Plates)
Set 1: 30 kg x 10
Set 2: 30 kg x 10
Set 3: 30 kg x 10
Set 4: 30 kg x 7
Set 5: 30 kg x 7

Triceps Extension (Machine)
Set 1: 50 kg x 8
Set 2: 50 kg x 8
Set 3: 50 kg x 8
Set 4: 50 kg x 8

Bicep Curl (Machine)
Set 1: 30 kg x 10
Set 2: 30 kg x 10
Set 3: 30 kg x 10

Lateral Raise (Cable)
Set 1: 5 kg x 8
Set 2: 5 kg x 8
Set 3: 5 kg x 8

Triceps Extension (Cable)
Set 1: 23 kg x 10
Set 2: 23 kg x 10
Set 3: 23 kg x 10

@hevyapp
https://hevy.com/workout/OrKJRBEijwm"""


def test_hevy_shared_text_is_read_set_by_set():
    sets, unparsed = workout_text.parse(HEVY_SHARE)
    assert unparsed == []
    assert len(sets) == 22
    dl = [(s.weight_kg, s.reps) for s in sets if s.exercise == "sumo deadlift"]
    assert dl == [(60.0, 5), (90.0, 5), (90.0, 5), (90.0, 5)]
    sp = [s.reps for s in sets if s.exercise == "shoulder press (machine plates)"]
    assert sp == [10, 10, 10, 7, 7] and [s.set_no for s in sets if s.exercise == "sumo deadlift"] == [1, 2, 3, 4]


def test_changing_reps_on_one_line():
    sets, unparsed = workout_text.parse("Shoulder press 3x10 + 2x7 30kg\nPlank 2x45s")
    assert unparsed == []
    assert [(s.reps, s.weight_kg) for s in sets[:5]] == [(10, 30.0)] * 3 + [(7, 30.0)] * 2
    assert sets[-1].duration_s == 45
    sets, _ = workout_text.parse("Squat 1x5 60kg, 3x5 90kg")
    assert [s.weight_kg for s in sets] == [60.0, 90.0, 90.0, 90.0]


def test_unknown_lines_are_reported_not_guessed():
    sets, unparsed = workout_text.parse("Squat 3x5 80kg\nHeavy day, felt great!\nSet 1: 50 kg x 5")
    assert len(sets) == 3 and "Heavy day, felt great!" in unparsed and "Set 1: 50 kg x 5" in unparsed


def test_tools_state_today(env):
    from app import plans
    conn, runner, _, _ = env
    s = plans.sessions_around(conn, runner["runner_id"], date(2026, 10, 24))
    assert s["today_is"] == {"date": "2026-10-24", "weekday": "Saturday", "tomorrow": "Sunday 2026-10-25"}
